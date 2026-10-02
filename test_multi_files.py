import io
import docx
import fitz
from PIL import Image, ImageDraw
from fastapi.testclient import TestClient
from app.main import app

def test_multifile():
    client = TestClient(app)

    # 1. Create a Word document in memory
    doc = docx.Document()
    doc.add_heading('Giao an Toan 12: Ung dung tich phan', level=1)
    doc.add_paragraph('Bai 1: Tinh dien tich hinh phang gioi han boi y = x^2 va y = 2x.')
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_bytes = doc_io.getvalue()

    # 2. Create a PDF in memory
    pdf_doc = fitz.open()
    page = pdf_doc.new_page()
    page.insert_text((50, 72), 'De cuong on tap kiem tra hoc ky: Mon Toan Lop 12')
    page.insert_text((50, 100), 'Phan 1: Tich phan xac dinh va phuong phap doi bien.')
    pdf_bytes = pdf_doc.tobytes()

    # 3. Create an Image in memory
    img = Image.new('RGB', (350, 80), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((10, 20), 'Bai tap hinh hoc khong gian Oxyz:', fill=(0, 0, 0))
    d.text((10, 45), 'Cho mat cau (S): x^2 + y^2 + z^2 - 4x + 2y - 4 = 0', fill=(0, 0, 0))
    img_io = io.BytesIO()
    img.save(img_io, format='PNG')
    img_bytes = img_io.getvalue()

    # 4. Multi-file upload (Word + PDF + Image)
    files = [
        ('files', ('giao_an_toan_12.docx', doc_bytes, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')),
        ('files', ('de_cuong_hoc_ky.pdf', pdf_bytes, 'application/pdf')),
        ('files', ('bai_tap_anh.png', img_bytes, 'image/png'))
    ]

    print('Sending multi-file extract request (Word + PDF + Image)...')
    r = client.post('/api/extract', files=files)
    assert r.status_code == 200
    data = r.json()
    assert data['total_files'] == 3
    assert len(data['files']) == 3
    
    print(f"[PASS] Total files processed: {data['total_files']}")
    for f in data['files']:
        print(f"  - [{f['category']}] {f['filename']}: {f['size_formatted']}, {f['length']} ký tự, status={f['status']}")
        assert f['status'] == 'success'
        assert f['length'] > 0
        
    assert 'Ung dung tich phan' in data['combined_text']
    assert 'De cuong on tap kiem tra' in data['combined_text']
    assert 'mat cau' in data['combined_text'].lower() or 'hinh hoc' in data['combined_text'].lower() or 'bai tap' in data['combined_text'].lower() or 'tệp ảnh' in data['combined_text'].lower()
    
    print("\n[PASS] Multi-file extraction (Word + PDF + Image) succeeded 100%!")

if __name__ == '__main__':
    test_multifile()
