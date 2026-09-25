import sys
import io
import docx
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

from app.services.ai_generator import get_mock_math_exam
from app.services.docx_exporter import create_exam_document

def test_red_answers_formatting():
    print("=== KIỂM THỬ ĐỊNH DẠNG ĐÁP ÁN IN ĐỎ (WORD EXPORT) ===")
    
    exam = get_mock_math_exam()
    # Thêm 1 câu tự luận mẫu để kiểm thử Phần IV
    from app.services.models import Part4EssayQuestion
    exam.part4_essay = [
        Part4EssayQuestion(
            id=1,
            question="Cho hình chóp $S.ABC$ có đáy là tam giác đều cạnh $a$. Tính thể tích.",
            points=1.0,
            explanation="- Diện tích đáy: $S = \\frac{a^2\\sqrt{3}}{4}$.\n- Chiều cao: $h = a\\sqrt{3}$.\n- Thể tích: $V = \\frac{a^3}{4}$."
        )
    ]
    
    doc = create_exam_document(exam, variant_code="101", red_answers=True)
    
    # Quét tất cả các đoạn văn (paragraphs) và bảng (tables) để tìm các run màu đỏ
    red_runs = []
    normal_runs = []
    
    def check_paragraph(p, context=""):
        for r in p.runs:
            c = r.font.color
            is_red = False
            if c and c.rgb:
                # RGBColor(192, 0, 0)
                if c.rgb[0] > 150 and c.rgb[1] < 50 and c.rgb[2] < 50:
                    is_red = True
            if is_red:
                red_runs.append((context, r.text, r.bold, r.italic))
            else:
                normal_runs.append((context, r.text, r.bold, r.italic))
                
        # Kiểm tra OMML math elements trong paragraph oxml
        for math_node in p._p.xpath('.//*[local-name()="oMath" or local-name()="oMathPara"]'):
            for r_node in math_node.xpath('.//*[local-name()="r"]'):
                # Kiểm tra color
                rPr = r_node.xpath('.//*[local-name()="rPr"]')
                if rPr:
                    color_node = rPr[0].xpath('.//*[local-name()="color"]')
                    bold_node = rPr[0].xpath('.//*[local-name()="b"]')
                    i_node = rPr[0].xpath('.//*[local-name()="i"]')
                    txt_nodes = r_node.xpath('.//*[local-name()="t"]')
                    m_txt = "".join([t.text for t in txt_nodes if t.text])
                    if color_node:
                        val = color_node[0].get(qn('w:val'), '').upper()
                        if val in ('C00000', 'FF0000', 'DC2626'):
                            has_bold = len(bold_node) > 0
                            has_i = len(i_node) > 0
                            red_runs.append((f"{context} (OMML)", m_txt, has_bold, has_i))

    for p in doc.paragraphs:
        check_paragraph(p, "Body Paragraph")
        
    for t in doc.tables[1:]: # Bỏ qua header_table ở đầu trang
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    check_paragraph(p, "Table Cell")

    print(f"Tổng số phần tử đáp án in đỏ tìm thấy: {len(red_runs)}")
    assert len(red_runs) > 0, "Không tìm thấy đáp án in đỏ nào!"
    
    errors = []
    for ctx, txt, bold, italic in red_runs:
        if bold:
            errors.append(f"[LỖI IN ĐẬM]: '{txt}' tại {ctx} bị in đậm (bold=True)!")
        if italic:
            errors.append(f"[LỖI IN NGHIÊNG]: '{txt}' tại {ctx} bị in nghiêng (italic=True)!")
            
    if errors:
        print(f"❌ Phát hiện {len(errors)} lỗi định dạng:")
        for e in errors[:10]:
            print("  ", e)
        raise AssertionError("Có đáp án màu đỏ vẫn bị in đậm hoặc in nghiêng!")
        
    print("✅ XÁC NHẬN: 100% đáp án in đỏ KHÔNG BỊ IN ĐẬM VÀ KHÔNG BỊ IN NGHIÊNG!")
    print(f"   -> Đã kiểm tra thành công {len(red_runs)} phần tử màu đỏ trên toàn bộ đề thi.")

if __name__ == "__main__":
    test_red_answers_formatting()
