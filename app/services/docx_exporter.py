import os
import re
import io
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import latex2mathml.converter
from lxml import etree
from typing import List, Optional, Dict, Any

from .models import ExamStructure, ExamVariant, Option, Part1Question, Part2Question, Part3Question, ExamScoring, calculate_exam_scoring, sync_part4_essay_points, clean_essay_explanation
from .explanation_sync import synchronize_mcq_explanation_with_answer

# Locate MML2OMML.XSL
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.dirname(CURRENT_DIR)
XSL_PATH = os.path.join(PROJECT_DIR, 'templates', 'MML2OMML.XSL')
if not os.path.exists(XSL_PATH):
    # Fallback to Office system path
    system_xsl = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
    if os.path.exists(system_xsl):
        XSL_PATH = system_xsl

_xslt_transform = None

def get_xslt_transform():
    global _xslt_transform
    if _xslt_transform is None and os.path.exists(XSL_PATH):
        try:
            xslt_tree = etree.parse(XSL_PATH)
            _xslt_transform = etree.XSLT(xslt_tree)
        except Exception as e:
            print("Error loading XSLT:", e)
    return _xslt_transform

def latex_to_omml(latex_code: str):
    """Convert a LaTeX formula string to an lxml OMML element."""
    transform = get_xslt_transform()
    if transform is None:
        return None
    try:
        # Clean latex string
        clean_latex = latex_code.strip()
        # Clean common problematic tokens
        clean_latex = clean_latex.replace(r'\,', ' ').replace(r'\;', ' ').replace(r'\quad', ' ')
        mathml = latex2mathml.converter.convert(clean_latex)
        mml_tree = etree.fromstring(mathml)
        omml = transform(mml_tree)
        return omml.getroot()
    except Exception as e:
        # If conversion fails, return None to fallback
        return None

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set table cell padding in twips (1/20 of a pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """Set cell borders: top, bottom, left, right with e.g. sz='6', val='single', color='CCCCCC'."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        if border_name in kwargs:
            edge = OxmlElement(f'w:{border_name}')
            for key, val in kwargs[border_name].items():
                edge.set(qn(f'w:{key}'), str(val))
            tcBorders.append(edge)
    tcPr.append(tcBorders)

def add_formatted_text_with_math(paragraph, text: str, bold=False, italic=False, font_name="Times New Roman", font_size=12, color=None):
    """
    Parses text containing $...$ or $$...$$ LaTeX formulas and inserts
    regular text runs and native Word OMML equations.
    """
    if not text:
        return
    
    # Split text by math delimiters: $$...$$ or $...$
    pattern = r'(\$\$.*?\$\$|\$.*?\$)'
    tokens = re.split(pattern, text)
    
    for token in tokens:
        if not token:
            continue
        if (token.startswith('$$') and token.endswith('$$') and len(token) > 4) or \
           (token.startswith('$') and token.endswith('$') and len(token) > 2):
            # Extract raw formula
            is_display = token.startswith('$$')
            raw_math = token[2:-2] if is_display else token[1:-1]
            omml_elem = latex_to_omml(raw_math)
            if omml_elem is not None:
                # Convert lxml OMML element into docx oxml element
                xml_str = etree.tostring(omml_elem, encoding='utf-8')
                docx_math_element = parse_xml(xml_str)
                if color:
                    hex_color = f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"
                    for r_node in docx_math_element.xpath('.//*[local-name()="r"]'):
                        rPr = parse_xml(f'<w:rPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:color w:val="{hex_color}"/>{"<w:b/>" if bold else ""}</w:rPr>')
                        r_node.insert(0, rPr)
                paragraph._p.append(docx_math_element)
            else:
                # Fallback to plain italic run
                run = paragraph.add_run(raw_math)
                run.bold = bold
                run.italic = True
                run.font.name = font_name
                run.font.size = Pt(font_size)
                if color:
                    run.font.color.rgb = color
        else:
            # Regular text
            run = paragraph.add_run(token)
            run.bold = bold
            run.italic = italic
            run.font.name = font_name
            run.font.size = Pt(font_size)
            if color:
                run.font.color.rgb = color

def format_points(val: Any, decimals: int = 1) -> str:
    try:
        f_val = float(val)
        if f_val.is_integer():
            return str(int(f_val))
        return f"{f_val:.{decimals}f}".replace(".", ",")
    except Exception:
        return str(val)

def setup_exam_footer(section, code: str):
    """
    Thiết lập chân trang chuẩn cho đề thi:
    - Bên trái: Mã đề: {code}
    - Bên phải: Trang {PAGE}/{NUMPAGES} (tự động cập nhật cho đến hết)
    """
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.77), WD_TAB_ALIGNMENT.RIGHT)
    
    # Bên trái: Mã đề
    display_code = "Tổng hợp" if str(code).upper() == "TONG_HOP" else str(code)
    r_code = p.add_run(f"Mã đề: {display_code}")
    r_code.font.name = "Times New Roman"
    r_code.font.size = Pt(9.5)
    r_code.font.italic = True
    r_code.font.color.rgb = RGBColor(100, 100, 100)
    
    # Tab sang sát lề phải
    p.add_run("\t")
    
    # Bên phải: Trang PAGE/NUMPAGES
    r_pfx = p.add_run("Trang ")
    r_pfx.font.name = "Times New Roman"
    r_pfx.font.size = Pt(9.5)
    r_pfx.font.color.rgb = RGBColor(100, 100, 100)
    
    fld_page = parse_xml(r'<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:instr="PAGE"><w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="19"/><w:color w:val="646464"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>')
    p._p.append(fld_page)
    
    r_slash = p.add_run("/")
    r_slash.font.name = "Times New Roman"
    r_slash.font.size = Pt(9.5)
    r_slash.font.color.rgb = RGBColor(100, 100, 100)
    
    fld_numpages = parse_xml(r'<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:instr="NUMPAGES"><w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="19"/><w:color w:val="646464"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>')
    p._p.append(fld_numpages)

def create_exam_document(
    exam: ExamStructure,
    variant_code: Optional[str] = None,
    include_answers: bool = True,
    include_explanations: bool = True,
    all_variants: Optional[List[ExamVariant]] = None,
    red_answers: bool = False
) -> docx.Document:
    """
    Generates a beautifully formatted Vietnamese curriculum exam paper (.docx).
    Structure:
      Section 1: ĐỀ THI (Header, Student Box, Part 1, Part 2, Part 3)
      Section 2: BẢNG ĐÁP ÁN (Matrix answers across codes or detailed table)
      Section 3: LỜI GIẢI CHI TIẾT (Step by step explanations with math)
    """
    doc = docx.Document()
    
    # Configure page setup: Standard A4, 2cm margins
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # A4 width
    section.page_height = Inches(11.69) # A4 height
    section.top_margin = Inches(0.75)   # ~1.9 cm
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    
    code = variant_code or exam.code or "101"
    setup_exam_footer(section, code)
    
    scoring = exam.scoring
    if scoring is None:
        p4_total = sum(q.points for q in (exam.part4_essay or []))
        scoring = calculate_exam_scoring(
            num_p1=len(exam.part1_mcq),
            num_p2=len(exam.part2_tf),
            num_p3=len(exam.part3_short),
            num_p4=len(exam.part4_essay or []),
            p4_points_total=p4_total
        )

    # Đảm bảo điểm các câu con tự luận luôn luôn bằng đúng điểm của Phần IV
    if exam.part4_essay and scoring:
        sync_part4_essay_points(exam.part4_essay, scoring.part4_points)
    
    # -------------------------------------------------------------
    # 1. HEADER (Two columns table: School Info | Exam & Test Code)
    # -------------------------------------------------------------
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(3.6)
    header_table.columns[1].width = Inches(3.4)
    
    cell_left = header_table.cell(0, 0)
    cell_right = header_table.cell(0, 1)
    
    # Left Header: School & Department
    p_left1 = cell_left.paragraphs[0]
    p_left1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_left1.paragraph_format.space_after = Pt(2)
    run_sch = p_left1.add_run(exam.school_name.upper() + "\n")
    run_sch.bold = True
    run_sch.font.name = "Times New Roman"
    run_sch.font.size = Pt(11)
    
    run_yr = p_left1.add_run(exam.academic_year)
    run_yr.italic = True
    run_yr.font.name = "Times New Roman"
    run_yr.font.size = Pt(10)
    
    # Right Header: Exam Title & Test Code
    p_right1 = cell_right.paragraphs[0]
    p_right1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_right1.paragraph_format.space_after = Pt(2)
    run_title = p_right1.add_run(f"{exam.title.upper()}\n")
    run_title.bold = True
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(11)
    
    run_sub = p_right1.add_run(f"Môn: {exam.subject.upper()} - Lớp: {exam.grade}\n")
    run_sub.bold = True
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(10.5)
    
    run_dur = p_right1.add_run(f"Thời gian làm bài: {exam.duration_minutes} phút\n(Không kể thời gian phát đề)\n")
    run_dur.italic = True
    run_dur.font.name = "Times New Roman"
    run_dur.font.size = Pt(9.5)
    
    # Test Code Highlight Box
    p_code = cell_right.add_paragraph()
    p_code.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_code.paragraph_format.space_before = Pt(4)
    p_code.paragraph_format.space_after = Pt(2)
    run_code_lbl = p_code.add_run("MÃ ĐỀ THI: ")
    run_code_lbl.bold = True
    run_code_lbl.font.name = "Times New Roman"
    run_code_lbl.font.size = Pt(11)
    run_code_val = p_code.add_run(f" {code} ")
    run_code_val.bold = True
    run_code_val.font.name = "Times New Roman"
    run_code_val.font.size = Pt(12)
    run_code_val.font.color.rgb = RGBColor(180, 0, 0)
    
    # Student info box
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    info_p = doc.add_paragraph()
    info_p.paragraph_format.space_before = Pt(4)
    info_p.paragraph_format.space_after = Pt(8)
    run_info = info_p.add_run("Họ và tên thí sinh: ............................................................................ Số báo danh: .............................")
    run_info.italic = True
    run_info.font.name = "Times New Roman"
    run_info.font.size = Pt(10.5)
    
    # Divider line
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_before = Pt(0)
    div_p.paragraph_format.space_after = Pt(12)
    div_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_div = div_p.add_run("━" * 50)
    run_div.font.color.rgb = RGBColor(150, 150, 150)
    run_div.font.size = Pt(8)
    
    # -------------------------------------------------------------
    # SECTION 1: ĐỀ THI - PHẦN I (TRẮC NGHIỆM NHIỀU LỰA CHỌN)
    # -------------------------------------------------------------
    if exam.part1_mcq:
        p_p1_title = doc.add_paragraph()
        p_p1_title.paragraph_format.space_before = Pt(8)
        p_p1_title.paragraph_format.space_after = Pt(4)
        run_p1 = p_p1_title.add_run(f"PHẦN I ({format_points(scoring.part1_points)} điểm). Thí sinh trả lời từ câu 1 đến câu {len(exam.part1_mcq)}. Mỗi câu hỏi thí sinh chỉ chọn một phương án. (Mỗi câu trả lời đúng được {format_points(scoring.part1_per_q, 2)} điểm)")
        run_p1.bold = True
        run_p1.font.name = "Times New Roman"
        run_p1.font.size = Pt(11)
        run_p1.font.color.rgb = RGBColor(0, 51, 102)
        
        for idx, q in enumerate(exam.part1_mcq, start=1):
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(6)
            p_q.paragraph_format.space_after = Pt(3)
            p_q.paragraph_format.line_spacing = 1.15
            
            # Question number
            run_num = p_q.add_run(f"Câu {idx}: ")
            run_num.bold = True
            run_num.font.name = "Times New Roman"
            run_num.font.size = Pt(11)
            
            # Question body with math
            add_formatted_text_with_math(p_q, q.question, font_size=11)
            
            red_color = RGBColor(192, 0, 0)
            
            # Options formatting (Layout neatly according to text length)
            opts = q.options[:4]
            total_opt_len = sum(len(opt.text) for opt in opts)
            
            if total_opt_len > 120 or len(opts) < 4:
                # 1 column (vertical paragraphs)
                for opt in opts:
                    p_opt = doc.add_paragraph()
                    p_opt.paragraph_format.left_indent = Inches(0.25)
                    p_opt.paragraph_format.space_before = Pt(1)
                    p_opt.paragraph_format.space_after = Pt(1)
                    
                    is_correct = (red_answers and opt.label.strip().upper() == q.answer.strip().upper())
                    opt_color = red_color if is_correct else None
                    opt_bold = True if is_correct else False
                    
                    run_lbl = p_opt.add_run(f"{opt.label}. ")
                    run_lbl.bold = True
                    run_lbl.font.name = "Times New Roman"
                    run_lbl.font.size = Pt(11)
                    if is_correct:
                        run_lbl.font.color.rgb = red_color
                    add_formatted_text_with_math(p_opt, opt.text, bold=opt_bold, font_size=11, color=opt_color)
            elif total_opt_len > 60:
                # 2 columns (2x2 table)
                t2 = doc.add_table(rows=2, cols=2)
                t2.alignment = WD_TABLE_ALIGNMENT.CENTER
                t2.autofit = True
                t2.columns[0].width = Inches(3.5)
                t2.columns[1].width = Inches(3.5)
                for opt_idx, opt in enumerate(opts):
                    r_idx = opt_idx // 2
                    c_idx = opt_idx % 2
                    cell = t2.cell(r_idx, c_idx)
                    p_cell = cell.paragraphs[0]
                    p_cell.paragraph_format.space_before = Pt(1)
                    p_cell.paragraph_format.space_after = Pt(1)
                    
                    is_correct = (red_answers and opt.label.strip().upper() == q.answer.strip().upper())
                    opt_color = red_color if is_correct else None
                    opt_bold = True if is_correct else False
                    
                    run_lbl = p_cell.add_run(f"{opt.label}. ")
                    run_lbl.bold = True
                    run_lbl.font.name = "Times New Roman"
                    run_lbl.font.size = Pt(11)
                    if is_correct:
                        run_lbl.font.color.rgb = red_color
                    add_formatted_text_with_math(p_cell, opt.text, bold=opt_bold, font_size=11, color=opt_color)
            else:
                # 4 columns (1x4 table)
                opt_table = doc.add_table(rows=1, cols=4)
                opt_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                opt_table.autofit = True
                for opt_idx, opt in enumerate(opts):
                    cell = opt_table.cell(0, opt_idx)
                    cell.width = Inches(1.75)
                    p_cell = cell.paragraphs[0]
                    p_cell.paragraph_format.space_before = Pt(1)
                    p_cell.paragraph_format.space_after = Pt(1)
                    
                    is_correct = (red_answers and opt.label.strip().upper() == q.answer.strip().upper())
                    opt_color = red_color if is_correct else None
                    opt_bold = True if is_correct else False
                    
                    run_lbl = p_cell.add_run(f"{opt.label}. ")
                    run_lbl.bold = True
                    run_lbl.font.name = "Times New Roman"
                    run_lbl.font.size = Pt(11)
                    if is_correct:
                        run_lbl.font.color.rgb = red_color
                    add_formatted_text_with_math(p_cell, opt.text, bold=opt_bold, font_size=11, color=opt_color)
    
    # -------------------------------------------------------------
    # SECTION 1: ĐỀ THI - PHẦN II (TRẮC NGHIỆM ĐÚNG SAI)
    # -------------------------------------------------------------
    if exam.part2_tf:
        p_p2_title = doc.add_paragraph()
        p_p2_title.paragraph_format.space_before = Pt(14)
        p_p2_title.paragraph_format.space_after = Pt(4)
        run_p2 = p_p2_title.add_run(f"PHẦN II ({format_points(scoring.part2_points)} điểm). Thí sinh trả lời từ câu 1 đến câu {len(exam.part2_tf)}. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.\n(Điểm tối đa mỗi câu là {format_points(scoring.part2_per_q)} điểm: Đúng 1 ý được 0,1 điểm; đúng 2 ý được 0,25 điểm; đúng 3 ý được 0,5 điểm; đúng 4 ý được 1,0 điểm)")
        run_p2.bold = True
        run_p2.font.name = "Times New Roman"
        run_p2.font.size = Pt(11)
        run_p2.font.color.rgb = RGBColor(0, 51, 102)
        
        for idx, q in enumerate(exam.part2_tf, start=1):
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(6)
            p_q.paragraph_format.space_after = Pt(3)
            p_q.paragraph_format.line_spacing = 1.15
            
            run_num = p_q.add_run(f"Câu {idx}: ")
            run_num.bold = True
            run_num.font.name = "Times New Roman"
            run_num.font.size = Pt(11)
            add_formatted_text_with_math(p_q, q.question, font_size=11)
            
            # Sub-items a, b, c, d
            for sub in q.sub_items:
                p_sub = doc.add_paragraph()
                p_sub.paragraph_format.left_indent = Inches(0.25)
                p_sub.paragraph_format.space_before = Pt(1)
                p_sub.paragraph_format.space_after = Pt(1)
                
                is_sub_correct = (red_answers and sub.is_correct is True)
                sub_color = red_color if is_sub_correct else None
                sub_bold = True if is_sub_correct else False
                
                run_lbl = p_sub.add_run(f"{sub.label}) ")
                run_lbl.bold = True
                run_lbl.font.name = "Times New Roman"
                run_lbl.font.size = Pt(11)
                if is_sub_correct:
                    run_lbl.font.color.rgb = red_color
                add_formatted_text_with_math(p_sub, sub.statement, bold=sub_bold, font_size=11, color=sub_color)
    
    # -------------------------------------------------------------
    # SECTION 1: ĐỀ THI - PHẦN III (TRẢ LỜI NGẮN)
    # -------------------------------------------------------------
    if exam.part3_short:
        p_p3_title = doc.add_paragraph()
        p_p3_title.paragraph_format.space_before = Pt(14)
        p_p3_title.paragraph_format.space_after = Pt(4)
        run_p3 = p_p3_title.add_run(f"PHẦN III ({format_points(scoring.part3_points)} điểm). Thí sinh trả lời từ câu 1 đến câu {len(exam.part3_short)}. Viết câu trả lời vào phiếu trả lời theo mẫu quy định. (Mỗi câu trả lời đúng được {format_points(scoring.part3_per_q, 2)} điểm)")
        run_p3.bold = True
        run_p3.font.name = "Times New Roman"
        run_p3.font.size = Pt(11)
        run_p3.font.color.rgb = RGBColor(0, 51, 102)
        
        for idx, q in enumerate(exam.part3_short, start=1):
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(6)
            p_q.paragraph_format.space_after = Pt(2 if red_answers else 4)
            p_q.paragraph_format.line_spacing = 1.15
            
            run_num = p_q.add_run(f"Câu {idx}: ")
            run_num.bold = True
            run_num.font.name = "Times New Roman"
            run_num.font.size = Pt(11)
            add_formatted_text_with_math(p_q, q.question, font_size=11)
            
            if red_answers and q.answer:
                p_ans = doc.add_paragraph()
                p_ans.paragraph_format.left_indent = Inches(0.25)
                p_ans.paragraph_format.space_before = Pt(1)
                p_ans.paragraph_format.space_after = Pt(4)
                p_ans.paragraph_format.line_spacing = 1.15
                r_ans_lbl = p_ans.add_run("Đáp án: ")
                r_ans_lbl.bold = True
                r_ans_lbl.font.name = "Times New Roman"
                r_ans_lbl.font.size = Pt(11)
                add_formatted_text_with_math(p_ans, q.answer, bold=True, font_size=11, color=red_color)
    
    # -------------------------------------------------------------
    # SECTION 1: ĐỀ THI - PHẦN IV (CÂU HỎI TỰ LUẬN)
    # -------------------------------------------------------------
    if exam.part4_essay and len(exam.part4_essay) > 0:
        p_p4_title = doc.add_paragraph()
        p_p4_title.paragraph_format.space_before = Pt(14)
        p_p4_title.paragraph_format.space_after = Pt(4)
        run_p4 = p_p4_title.add_run(f"PHẦN IV ({format_points(scoring.part4_points)} điểm). CÂU HỎI TỰ LUẬN (Thí sinh trình bày bài làm chi tiết vào tờ giấy thi)")
        run_p4.bold = True
        run_p4.font.name = "Times New Roman"
        run_p4.font.size = Pt(11)
        run_p4.font.color.rgb = RGBColor(0, 51, 102)
        
        for idx, q in enumerate(exam.part4_essay, start=1):
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(6)
            p_q.paragraph_format.space_after = Pt(2 if red_answers else 4)
            p_q.paragraph_format.line_spacing = 1.15
            
            pts_str = f" ({format_points(q.points)} điểm)" if q.points else ""
            run_num = p_q.add_run(f"Câu {idx}{pts_str}: ")
            run_num.bold = True
            run_num.font.name = "Times New Roman"
            run_num.font.size = Pt(11)
            add_formatted_text_with_math(p_q, q.question, font_size=11)
            
            if red_answers and (q.explanation or q.answer):
                p_hd = doc.add_paragraph()
                p_hd.paragraph_format.left_indent = Inches(0.25)
                p_hd.paragraph_format.space_before = Pt(2)
                p_hd.paragraph_format.space_after = Pt(1)
                r_hd = p_hd.add_run("Hướng dẫn chấm:")
                r_hd.italic = True
                r_hd.font.name = "Times New Roman"
                r_hd.font.size = Pt(10.5)
                
                exp_text = clean_essay_explanation(q.explanation or q.answer)
                for line in exp_text.split('\n'):
                    line_str = line.strip()
                    if line_str:
                        p_bullet = doc.add_paragraph()
                        p_bullet.paragraph_format.left_indent = Inches(0.35)
                        p_bullet.paragraph_format.space_before = Pt(1)
                        p_bullet.paragraph_format.space_after = Pt(1)
                        add_formatted_text_with_math(p_bullet, line_str, bold=False, font_size=10.5, color=red_color)
    
    # End of exam marker
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(18)
    p_end.paragraph_format.space_after = Pt(8)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_end = p_end.add_run("---------- HẾT ----------")
    run_end.bold = True
    run_end.font.name = "Times New Roman"
    run_end.font.size = Pt(11)
    
    if red_answers:
        return doc
    
    # -------------------------------------------------------------
    # SECTION 2: BẢNG ĐÁP ÁN (NEW PAGE)
    # -------------------------------------------------------------
    if include_answers:
        doc.add_page_break()
        
        p_ans_head = doc.add_paragraph()
        p_ans_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ans_head.paragraph_format.space_before = Pt(6)
        p_ans_head.paragraph_format.space_after = Pt(12)
        r_head = p_ans_head.add_run("BẢNG ĐÁP ÁN VÀ MA TRẬN ĐÁP ÁN\n")
        r_head.bold = True
        r_head.font.name = "Times New Roman"
        r_head.font.size = Pt(14)
        r_head.font.color.rgb = RGBColor(180, 0, 0)
        
        r_subhead = p_ans_head.add_run(f"Môn: {exam.subject} - Lớp: {exam.grade} | Mã đề: {code}")
        r_subhead.bold = True
        r_subhead.font.name = "Times New Roman"
        r_subhead.font.size = Pt(11)
        
        # 1. Part 1 Answer Table (20 MCQs)
        if exam.part1_mcq:
            p_t1 = doc.add_paragraph()
            p_t1.paragraph_format.space_before = Pt(8)
            p_t1.paragraph_format.space_after = Pt(4)
            r = p_t1.add_run(f"1. ĐÁP ÁN PHẦN I ({format_points(scoring.part1_points)} điểm - Mỗi câu đúng {format_points(scoring.part1_per_q, 2)} điểm)")
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0, 51, 102)
            
            # Generate table with 10 questions per row (2 rows of questions, 2 rows of answers)
            num_q = len(exam.part1_mcq)
            cols_per_row = 10
            
            for chunk_start in range(0, num_q, cols_per_row):
                chunk = exam.part1_mcq[chunk_start:chunk_start + cols_per_row]
                t_p1 = doc.add_table(rows=2, cols=len(chunk) + 1)
                t_p1.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                # Header col
                c00 = t_p1.cell(0, 0)
                c00.paragraphs[0].text = "Câu"
                c00.paragraphs[0].runs[0].font.bold = True
                c00.paragraphs[0].runs[0].font.name = "Times New Roman"
                c00.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                c00.width = Inches(0.8)
                set_cell_margins(c00)
                
                c10 = t_p1.cell(1, 0)
                c10.paragraphs[0].text = "Đ/Á"
                c10.paragraphs[0].runs[0].font.bold = True
                c10.paragraphs[0].runs[0].font.name = "Times New Roman"
                c10.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                c10.width = Inches(0.8)
                set_cell_margins(c10)
                
                for c_idx, q_item in enumerate(chunk):
                    # Question num
                    cell_q = t_p1.cell(0, c_idx + 1)
                    cell_q.paragraphs[0].text = str(chunk_start + c_idx + 1)
                    cell_q.paragraphs[0].runs[0].font.bold = True
                    cell_q.paragraphs[0].runs[0].font.name = "Times New Roman"
                    cell_q.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cell_q.width = Inches(0.6)
                    set_cell_margins(cell_q)
                    
                    # Answer letter
                    cell_a = t_p1.cell(1, c_idx + 1)
                    cell_a.paragraphs[0].text = q_item.answer.upper()
                    cell_a.paragraphs[0].runs[0].font.bold = True
                    cell_a.paragraphs[0].runs[0].font.name = "Times New Roman"
                    cell_a.paragraphs[0].runs[0].font.color.rgb = RGBColor(180, 0, 0)
                    cell_a.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cell_a.width = Inches(0.6)
                    set_cell_margins(cell_a)
                
                # Add border styling to table
                for row in t_p1.rows:
                    for cell in row.cells:
                        set_cell_border(cell,
                            top={'val': 'single', 'sz': '4', 'color': '888888'},
                            bottom={'val': 'single', 'sz': '4', 'color': '888888'},
                            left={'val': 'single', 'sz': '4', 'color': '888888'},
                            right={'val': 'single', 'sz': '4', 'color': '888888'}
                        )
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
        # 2. Part 2 Answer Table (True/False)
        if exam.part2_tf:
            p_t2 = doc.add_paragraph()
            p_t2.paragraph_format.space_before = Pt(8)
            p_t2.paragraph_format.space_after = Pt(4)
            r = p_t2.add_run(f"2. ĐÁP ÁN PHẦN II ({format_points(scoring.part2_points)} điểm - Tối đa {format_points(scoring.part2_per_q)} đ/câu: 1 ý=0,1đ; 2 ý=0,25đ; 3 ý=0,5đ; 4 ý=1,0đ)")
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0, 51, 102)
            
            t_p2 = doc.add_table(rows=len(exam.part2_tf) + 1, cols=5)
            t_p2.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            headers = ["Câu", "Lệnh a)", "Lệnh b)", "Lệnh c)", "Lệnh d)"]
            for col_idx, h in enumerate(headers):
                cell = t_p2.cell(0, col_idx)
                cell.paragraphs[0].text = h
                cell.paragraphs[0].runs[0].font.bold = True
                cell.paragraphs[0].runs[0].font.name = "Times New Roman"
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_margins(cell)
            
            for row_idx, q_item in enumerate(exam.part2_tf, start=1):
                cell_c = t_p2.cell(row_idx, 0)
                cell_c.paragraphs[0].text = f"Câu {row_idx}"
                cell_c.paragraphs[0].runs[0].font.bold = True
                cell_c.paragraphs[0].runs[0].font.name = "Times New Roman"
                cell_c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_margins(cell_c)
                
                # Map sub-items
                sub_dict = {s.label.lower(): ('Đ' if s.is_correct else 'S') for s in q_item.sub_items}
                for c_idx, lbl in enumerate(['a', 'b', 'c', 'd'], start=1):
                    val = sub_dict.get(lbl, '-')
                    cell_sub = t_p2.cell(row_idx, c_idx)
                    cell_sub.paragraphs[0].text = val
                    run_v = cell_sub.paragraphs[0].runs[0]
                    run_v.font.bold = True
                    run_v.font.name = "Times New Roman"
                    run_v.font.color.rgb = RGBColor(0, 130, 0) if val == 'Đ' else RGBColor(180, 0, 0)
                    cell_sub.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    set_cell_margins(cell_sub)
                    
            for row in t_p2.rows:
                for cell in row.cells:
                    set_cell_border(cell,
                        top={'val': 'single', 'sz': '4', 'color': '888888'},
                        bottom={'val': 'single', 'sz': '4', 'color': '888888'},
                        left={'val': 'single', 'sz': '4', 'color': '888888'},
                        right={'val': 'single', 'sz': '4', 'color': '888888'}
                    )
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
        # 3. Part 3 Answer Table (Short Answers)
        if exam.part3_short:
            p_t3 = doc.add_paragraph()
            p_t3.paragraph_format.space_before = Pt(8)
            p_t3.paragraph_format.space_after = Pt(4)
            r = p_t3.add_run(f"3. ĐÁP ÁN PHẦN III ({format_points(scoring.part3_points)} điểm - Mỗi câu đúng {format_points(scoring.part3_per_q, 2)} điểm)")
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0, 51, 102)
            
            t_p3 = doc.add_table(rows=2, cols=len(exam.part3_short) + 1)
            t_p3.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            c00 = t_p3.cell(0, 0)
            c00.paragraphs[0].text = "Câu"
            c00.paragraphs[0].runs[0].font.bold = True
            c00.paragraphs[0].runs[0].font.name = "Times New Roman"
            c00.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_margins(c00)
            
            c10 = t_p3.cell(1, 0)
            c10.paragraphs[0].text = "Đáp số"
            c10.paragraphs[0].runs[0].font.bold = True
            c10.paragraphs[0].runs[0].font.name = "Times New Roman"
            c10.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_margins(c10)
            
            for idx, q_item in enumerate(exam.part3_short, start=1):
                cq = t_p3.cell(0, idx)
                cq.paragraphs[0].text = str(idx)
                cq.paragraphs[0].runs[0].font.bold = True
                cq.paragraphs[0].runs[0].font.name = "Times New Roman"
                cq.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_margins(cq)
                
                ca = t_p3.cell(1, idx)
                ca.paragraphs[0].text = q_item.answer
                ca.paragraphs[0].runs[0].font.bold = True
                ca.paragraphs[0].runs[0].font.name = "Times New Roman"
                ca.paragraphs[0].runs[0].font.color.rgb = RGBColor(180, 0, 0)
                ca.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_margins(ca)
                
            for row in t_p3.rows:
                for cell in row.cells:
                    set_cell_border(cell,
                        top={'val': 'single', 'sz': '4', 'color': '888888'},
                        bottom={'val': 'single', 'sz': '4', 'color': '888888'},
                        left={'val': 'single', 'sz': '4', 'color': '888888'},
                        right={'val': 'single', 'sz': '4', 'color': '888888'}
                    )
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
        
        # 4. Part 4 Essay Guide (if present)
        if exam.part4_essay and len(exam.part4_essay) > 0:
            p_t4 = doc.add_paragraph()
            p_t4.paragraph_format.space_before = Pt(10)
            p_t4.paragraph_format.space_after = Pt(4)
            r = p_t4.add_run(f"4. HƯỚNG DẪN CHẤM PHẦN TỰ LUẬN ({format_points(scoring.part4_points)} điểm)")
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0, 51, 102)
            
            for idx, q in enumerate(exam.part4_essay, start=1):
                p_item = doc.add_paragraph()
                p_item.paragraph_format.space_before = Pt(4)
                p_item.paragraph_format.space_after = Pt(2)
                pts_str = f" ({format_points(q.points)} điểm)" if q.points else ""
                r_qnum = p_item.add_run(f"Câu {idx}{pts_str}: ")
                r_qnum.bold = True
                r_qnum.font.name = "Times New Roman"
                r_qnum.font.size = Pt(11)
                
                if q.answer:
                    r_ans = p_item.add_run(f"Đáp án then chốt: {q.answer}\n")
                    r_ans.font.name = "Times New Roman"
                    r_ans.font.size = Pt(10.5)
                
                if q.explanation:
                    p_rubric = doc.add_paragraph()
                    p_rubric.paragraph_format.left_indent = Inches(0.25)
                    p_rubric.paragraph_format.space_before = Pt(1)
                    p_rubric.paragraph_format.space_after = Pt(3)
                    add_formatted_text_with_math(p_rubric, clean_essay_explanation(q.explanation), font_size=10.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # 5. Master Comparison Table if multiple variants are supplied
        if all_variants and len(all_variants) > 1:
            p_mat = doc.add_paragraph()
            p_mat.paragraph_format.space_before = Pt(14)
            p_mat.paragraph_format.space_after = Pt(4)
            r_mat = p_mat.add_run(f"MA TRẬN ĐÁP ÁN TRẮC NGHIỆM PHẦN I GIỮA CÁC MÃ ĐỀ (TỔNG CỘNG {len(all_variants)} MÃ ĐỀ)")
            r_mat.bold = True
            r_mat.font.name = "Times New Roman"
            r_mat.font.size = Pt(11)
            r_mat.font.color.rgb = RGBColor(150, 0, 100)
            
            num_mcqs = len(all_variants[0].exam.part1_mcq)
            chunk_size = 8
            for chunk_start in range(0, len(all_variants), chunk_size):
                v_chunk = all_variants[chunk_start:chunk_start + chunk_size]
                chunk_codes_str = ", ".join(v.code for v in v_chunk)
                
                if len(all_variants) > chunk_size:
                    p_submat = doc.add_paragraph()
                    p_submat.paragraph_format.space_before = Pt(8)
                    p_submat.paragraph_format.space_after = Pt(2)
                    r_sm = p_submat.add_run(f"✦ Bảng đối chiếu nhóm mã đề: {chunk_codes_str}")
                    r_sm.bold = True
                    r_sm.font.name = "Times New Roman"
                    r_sm.font.size = Pt(10)
                    r_sm.font.color.rgb = RGBColor(0, 51, 102)
                
                mat_table = doc.add_table(rows=num_mcqs + 1, cols=len(v_chunk) + 1)
                mat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                # Header
                c0 = mat_table.cell(0, 0)
                c0.paragraphs[0].text = "Câu"
                c0.paragraphs[0].runs[0].font.bold = True
                c0.paragraphs[0].runs[0].font.name = "Times New Roman"
                c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_margins(c0)
                
                for v_idx, v in enumerate(v_chunk, start=1):
                    cv = mat_table.cell(0, v_idx)
                    cv.paragraphs[0].text = f"Mã {v.code}"
                    cv.paragraphs[0].runs[0].font.bold = True
                    cv.paragraphs[0].runs[0].font.name = "Times New Roman"
                    cv.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    set_cell_margins(cv)
                    
                for q_idx in range(num_mcqs):
                    row_c = mat_table.cell(q_idx + 1, 0)
                    row_c.paragraphs[0].text = str(q_idx + 1)
                    row_c.paragraphs[0].runs[0].font.bold = True
                    row_c.paragraphs[0].runs[0].font.name = "Times New Roman"
                    row_c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    set_cell_margins(row_c)
                    
                    for v_idx, v in enumerate(v_chunk, start=1):
                        ans_val = v.part1_answers.get(q_idx + 1, v.exam.part1_mcq[q_idx].answer if q_idx < len(v.exam.part1_mcq) else '-')
                        cell_a = mat_table.cell(q_idx + 1, v_idx)
                        cell_a.paragraphs[0].text = ans_val
                        cell_a.paragraphs[0].runs[0].font.bold = True
                        cell_a.paragraphs[0].runs[0].font.name = "Times New Roman"
                        cell_a.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        set_cell_margins(cell_a)
                        
                for row in mat_table.rows:
                    for cell in row.cells:
                        set_cell_border(cell,
                            top={'val': 'single', 'sz': '4', 'color': '888888'},
                            bottom={'val': 'single', 'sz': '4', 'color': '888888'},
                            left={'val': 'single', 'sz': '4', 'color': '888888'},
                            right={'val': 'single', 'sz': '4', 'color': '888888'}
                        )
                doc.add_paragraph().paragraph_format.space_after = Pt(4)

            # 5.2 Master Comparison Table for Part 2 (True / False)
            if all_variants[0].exam.part2_tf and len(all_variants[0].exam.part2_tf) > 0:
                p_mat2 = doc.add_paragraph()
                p_mat2.paragraph_format.space_before = Pt(12)
                p_mat2.paragraph_format.space_after = Pt(4)
                r_mat2 = p_mat2.add_run(f"MA TRẬN ĐÁP ÁN ĐÚNG / SAI PHẦN II GIỮA CÁC MÃ ĐỀ (TỔNG CỘNG {len(all_variants)} MÃ ĐỀ)")
                r_mat2.bold = True
                r_mat2.font.name = "Times New Roman"
                r_mat2.font.size = Pt(11)
                r_mat2.font.color.rgb = RGBColor(0, 102, 51)
                
                num_tf = len(all_variants[0].exam.part2_tf)
                sub_labels = ['a', 'b', 'c', 'd']
                total_tf_subitems = num_tf * 4
                
                for chunk_start in range(0, len(all_variants), chunk_size):
                    v_chunk = all_variants[chunk_start:chunk_start + chunk_size]
                    chunk_codes_str = ", ".join(v.code for v in v_chunk)
                    
                    if len(all_variants) > chunk_size:
                        p_submat2 = doc.add_paragraph()
                        p_submat2.paragraph_format.space_before = Pt(6)
                        p_submat2.paragraph_format.space_after = Pt(2)
                        r_sm2 = p_submat2.add_run(f"✦ Phần II - Nhóm mã đề: {chunk_codes_str}")
                        r_sm2.bold = True
                        r_sm2.font.name = "Times New Roman"
                        r_sm2.font.size = Pt(10)
                        r_sm2.font.color.rgb = RGBColor(0, 102, 51)
                    
                    mat2_table = doc.add_table(rows=total_tf_subitems + 1, cols=len(v_chunk) + 1)
                    mat2_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    
                    # Header
                    c0 = mat2_table.cell(0, 0)
                    c0.paragraphs[0].text = "Câu / Lệnh"
                    c0.paragraphs[0].runs[0].font.bold = True
                    c0.paragraphs[0].runs[0].font.name = "Times New Roman"
                    c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    set_cell_margins(c0)
                    
                    for v_idx, v in enumerate(v_chunk, start=1):
                        cv = mat2_table.cell(0, v_idx)
                        cv.paragraphs[0].text = f"Mã {v.code}"
                        cv.paragraphs[0].runs[0].font.bold = True
                        cv.paragraphs[0].runs[0].font.name = "Times New Roman"
                        cv.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        set_cell_margins(cv)
                    
                    # Rows for each question and subitem a, b, c, d
                    curr_row = 1
                    for q_idx in range(1, num_tf + 1):
                        for s_lbl in sub_labels:
                            row_c = mat2_table.cell(curr_row, 0)
                            row_c.paragraphs[0].text = f"Câu {q_idx}{s_lbl}"
                            row_c.paragraphs[0].runs[0].font.bold = True
                            row_c.paragraphs[0].runs[0].font.name = "Times New Roman"
                            row_c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                            set_cell_margins(row_c)
                            
                            for v_idx, v in enumerate(v_chunk, start=1):
                                cell_a = mat2_table.cell(curr_row, v_idx)
                                ans_val = v.part2_answers.get(q_idx, {}).get(s_lbl, '-')
                                cell_a.paragraphs[0].text = ans_val
                                if len(cell_a.paragraphs[0].runs) > 0:
                                    r_ans = cell_a.paragraphs[0].runs[0]
                                    r_ans.font.bold = True
                                    r_ans.font.name = "Times New Roman"
                                    r_ans.font.color.rgb = RGBColor(0, 130, 0) if ans_val == 'Đ' else (RGBColor(180, 0, 0) if ans_val == 'S' else RGBColor(100, 100, 100))
                                cell_a.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                                set_cell_margins(cell_a)
                            curr_row += 1
                    
                    for row in mat2_table.rows:
                        for cell in row.cells:
                            set_cell_border(cell,
                                top={'val': 'single', 'sz': '4', 'color': '888888'},
                                bottom={'val': 'single', 'sz': '4', 'color': '888888'},
                                left={'val': 'single', 'sz': '4', 'color': '888888'},
                                right={'val': 'single', 'sz': '4', 'color': '888888'}
                            )
                    doc.add_paragraph().paragraph_format.space_after = Pt(4)

            # 5.3 Master Comparison Table for Part 3 (Short Answers)
            if all_variants[0].exam.part3_short and len(all_variants[0].exam.part3_short) > 0:
                p_mat3 = doc.add_paragraph()
                p_mat3.paragraph_format.space_before = Pt(12)
                p_mat3.paragraph_format.space_after = Pt(4)
                r_mat3 = p_mat3.add_run(f"MA TRẬN ĐÁP ÁN TRẢ LỜI NGẮN PHẦN III GIỮA CÁC MÃ ĐỀ (TỔNG CỘNG {len(all_variants)} MÃ ĐỀ)")
                r_mat3.bold = True
                r_mat3.font.name = "Times New Roman"
                r_mat3.font.size = Pt(11)
                r_mat3.font.color.rgb = RGBColor(0, 51, 102)
                
                num_short = len(all_variants[0].exam.part3_short)
                for chunk_start in range(0, len(all_variants), chunk_size):
                    v_chunk = all_variants[chunk_start:chunk_start + chunk_size]
                    chunk_codes_str = ", ".join(v.code for v in v_chunk)
                    
                    if len(all_variants) > chunk_size:
                        p_submat3 = doc.add_paragraph()
                        p_submat3.paragraph_format.space_before = Pt(6)
                        p_submat3.paragraph_format.space_after = Pt(2)
                        r_sm3 = p_submat3.add_run(f"✦ Phần III - Nhóm mã đề: {chunk_codes_str}")
                        r_sm3.bold = True
                        r_sm3.font.name = "Times New Roman"
                        r_sm3.font.size = Pt(10)
                        r_sm3.font.color.rgb = RGBColor(0, 51, 102)
                    
                    mat3_table = doc.add_table(rows=num_short + 1, cols=len(v_chunk) + 1)
                    mat3_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    
                    # Header
                    c0 = mat3_table.cell(0, 0)
                    c0.paragraphs[0].text = "Câu"
                    c0.paragraphs[0].runs[0].font.bold = True
                    c0.paragraphs[0].runs[0].font.name = "Times New Roman"
                    c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    set_cell_margins(c0)
                    
                    for v_idx, v in enumerate(v_chunk, start=1):
                        cv = mat3_table.cell(0, v_idx)
                        cv.paragraphs[0].text = f"Mã {v.code}"
                        cv.paragraphs[0].runs[0].font.bold = True
                        cv.paragraphs[0].runs[0].font.name = "Times New Roman"
                        cv.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        set_cell_margins(cv)
                        
                    for q_idx in range(num_short):
                        row_c = mat3_table.cell(q_idx + 1, 0)
                        row_c.paragraphs[0].text = str(q_idx + 1)
                        row_c.paragraphs[0].runs[0].font.bold = True
                        row_c.paragraphs[0].runs[0].font.name = "Times New Roman"
                        row_c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        set_cell_margins(row_c)
                        
                        for v_idx, v in enumerate(v_chunk, start=1):
                            cell_a = mat3_table.cell(q_idx + 1, v_idx)
                            ans_val = v.part3_answers.get(q_idx + 1, v.exam.part3_short[q_idx].answer if q_idx < len(v.exam.part3_short) else '-')
                            cell_a.paragraphs[0].text = ans_val
                            if len(cell_a.paragraphs[0].runs) > 0:
                                r_ans = cell_a.paragraphs[0].runs[0]
                                r_ans.font.bold = True
                                r_ans.font.name = "Times New Roman"
                                r_ans.font.color.rgb = RGBColor(180, 0, 0)
                            cell_a.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                            set_cell_margins(cell_a)
                            
                    for row in mat3_table.rows:
                        for cell in row.cells:
                            set_cell_border(cell,
                                top={'val': 'single', 'sz': '4', 'color': '888888'},
                                bottom={'val': 'single', 'sz': '4', 'color': '888888'},
                                left={'val': 'single', 'sz': '4', 'color': '888888'},
                                right={'val': 'single', 'sz': '4', 'color': '888888'}
                            )
                    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # -------------------------------------------------------------
    # SECTION 3: HƯỚNG DẪN GIẢI CHI TIẾT (NEW PAGE)
    # -------------------------------------------------------------
    if include_explanations:
        doc.add_page_break()
        
        p_sol_head = doc.add_paragraph()
        p_sol_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sol_head.paragraph_format.space_before = Pt(6)
        p_sol_head.paragraph_format.space_after = Pt(12)
        r_head = p_sol_head.add_run("HƯỚNG DẪN GIẢI CHI TIẾT\n")
        r_head.bold = True
        r_head.font.name = "Times New Roman"
        r_head.font.size = Pt(14)
        r_head.font.color.rgb = RGBColor(0, 102, 51)
        
        r_subhead = p_sol_head.add_run(f"Môn: {exam.subject} - Lớp: {exam.grade} | Mã đề: {code}")
        r_subhead.bold = True
        r_subhead.font.name = "Times New Roman"
        r_subhead.font.size = Pt(11)
        
        # Solutions for Part 1
        if exam.part1_mcq:
            p_s1 = doc.add_paragraph()
            p_s1.paragraph_format.space_before = Pt(10)
            p_s1.paragraph_format.space_after = Pt(4)
            r = p_s1.add_run(f"PHẦN I. LỜI GIẢI CÂU TRẮC NGHIỆM NHIỀU LỰA CHỌN ({format_points(scoring.part1_points)} điểm)")
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0, 51, 102)
            
            for idx, q in enumerate(exam.part1_mcq, start=1):
                p_item = doc.add_paragraph()
                p_item.paragraph_format.space_before = Pt(6)
                p_item.paragraph_format.space_after = Pt(2)
                
                r_qnum = p_item.add_run(f"Câu {idx} (Chọn {q.answer}): ")
                r_qnum.bold = True
                r_qnum.font.name = "Times New Roman"
                r_qnum.font.size = Pt(11)
                r_qnum.font.color.rgb = RGBColor(180, 0, 0)
                
                p_exp = doc.add_paragraph()
                p_exp.paragraph_format.left_indent = Inches(0.2)
                p_exp.paragraph_format.space_before = Pt(1)
                p_exp.paragraph_format.space_after = Pt(4)
                
                explanation_text = synchronize_mcq_explanation_with_answer(q.explanation, q.answer) if q.explanation else f"Do đó chọn đáp án {q.answer}."
                add_formatted_text_with_math(p_exp, explanation_text, font_size=10.5)
        
        # Solutions for Part 2
        if exam.part2_tf:
            p_s2 = doc.add_paragraph()
            p_s2.paragraph_format.space_before = Pt(12)
            p_s2.paragraph_format.space_after = Pt(4)
            r = p_s2.add_run(f"PHẦN II. LỜI GIẢI CÂU TRẮC NGHIỆM ĐÚNG SAI ({format_points(scoring.part2_points)} điểm)")
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0, 51, 102)
            
            for idx, q in enumerate(exam.part2_tf, start=1):
                p_item = doc.add_paragraph()
                p_item.paragraph_format.space_before = Pt(6)
                p_item.paragraph_format.space_after = Pt(2)
                
                r_qnum = p_item.add_run(f"Câu {idx}: ")
                r_qnum.bold = True
                r_qnum.font.name = "Times New Roman"
                r_qnum.font.size = Pt(11)
                
                for sub in q.sub_items:
                    p_sub = doc.add_paragraph()
                    p_sub.paragraph_format.left_indent = Inches(0.25)
                    p_sub.paragraph_format.space_before = Pt(1)
                    p_sub.paragraph_format.space_after = Pt(2)
                    
                    status_text = "ĐÚNG" if sub.is_correct else "SAI"
                    r_lbl = p_sub.add_run(f"Ý {sub.label}) [{status_text}]: ")
                    r_lbl.bold = True
                    r_lbl.font.name = "Times New Roman"
                    r_lbl.font.size = Pt(10.5)
                    r_lbl.font.color.rgb = RGBColor(0, 130, 0) if sub.is_correct else RGBColor(180, 0, 0)
                    
                    sub_exp = sub.explanation or f"Khẳng định này là {'đúng' if sub.is_correct else 'sai'}."
                    add_formatted_text_with_math(p_sub, sub_exp, font_size=10.5)
        
        # Solutions for Part 3
        if exam.part3_short:
            p_s3 = doc.add_paragraph()
            p_s3.paragraph_format.space_before = Pt(12)
            p_s3.paragraph_format.space_after = Pt(4)
            r = p_s3.add_run(f"PHẦN III. LỜI GIẢI CÂU TRẢ LỜI NGẮN ({format_points(scoring.part3_points)} điểm)")
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0, 51, 102)
            
            for idx, q in enumerate(exam.part3_short, start=1):
                p_item = doc.add_paragraph()
                p_item.paragraph_format.space_before = Pt(6)
                p_item.paragraph_format.space_after = Pt(2)
                
                r_qnum = p_item.add_run(f"Câu {idx} (Đáp số: {q.answer}): ")
                r_qnum.bold = True
                r_qnum.font.name = "Times New Roman"
                r_qnum.font.size = Pt(11)
                r_qnum.font.color.rgb = RGBColor(180, 0, 0)
                
                p_exp = doc.add_paragraph()
                p_exp.paragraph_format.left_indent = Inches(0.2)
                p_exp.paragraph_format.space_before = Pt(1)
                p_exp.paragraph_format.space_after = Pt(4)
                
                exp_text = q.explanation or "Đang cập nhật lời giải..."
                add_formatted_text_with_math(p_exp, exp_text, font_size=10.5)

        # Solutions for Part 4 Essay
        if exam.part4_essay and len(exam.part4_essay) > 0:
            p_s4 = doc.add_paragraph()
            p_s4.paragraph_format.space_before = Pt(12)
            p_s4.paragraph_format.space_after = Pt(4)
            r = p_s4.add_run(f"PHẦN IV. LỜI GIẢI CHI TIẾT CÂU HỎI TỰ LUẬN ({format_points(scoring.part4_points)} điểm)")
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0, 51, 102)
            
            for idx, q in enumerate(exam.part4_essay, start=1):
                p_item = doc.add_paragraph()
                p_item.paragraph_format.space_before = Pt(6)
                p_item.paragraph_format.space_after = Pt(2)
                
                pts_str = f" ({format_points(q.points)} điểm)" if q.points else ""
                r_qnum = p_item.add_run(f"Câu {idx}{pts_str}: ")
                r_qnum.bold = True
                r_qnum.font.name = "Times New Roman"
                r_qnum.font.size = Pt(11)
                
                p_exp = doc.add_paragraph()
                p_exp.paragraph_format.left_indent = Inches(0.2)
                p_exp.paragraph_format.space_before = Pt(1)
                p_exp.paragraph_format.space_after = Pt(4)
                
                text_to_show = clean_essay_explanation(q.explanation) or clean_essay_explanation(q.answer) or "Học sinh giải và trình bày theo đúng các bước phương pháp."
                add_formatted_text_with_math(p_exp, text_to_show, font_size=10.5)

    return doc
