import io
import docx
import fitz  # PyMuPDF

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    full_text = []
    
    # 1. Đoạn văn thông thường
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            full_text.append(text)
    
    # 2. Bảng biểu (xử lý ô gộp merged cells và định dạng Markdown Table)
    if doc.tables:
        full_text.append("\n[BẢNG BIỂU / MA TRẬN TRÍCH XUẤT]:")
        for table_idx, table in enumerate(doc.tables, 1):
            table_lines = []
            for row in table.rows:
                seen_tc = set()
                row_vals = []
                for cell in row.cells:
                    if cell._tc in seen_tc:
                        continue
                    seen_tc.add(cell._tc)
                    clean_cell = " ".join(cell.text.split())
                    row_vals.append(clean_cell)
                if any(row_vals):
                    table_lines.append("| " + " | ".join(row_vals) + " |")
            if table_lines:
                full_text.append(f"\n--- Bảng {table_idx} ---")
                full_text.extend(table_lines)
                
    return '\n'.join(full_text)

def extract_text_from_pdf(file_bytes: bytes) -> str:
    doc = fitz.open(stream=file_bytes, filetype='pdf')
    full_text = []
    for page_idx, page in enumerate(doc, 1):
        # 1. Trích xuất bảng nếu có
        table_lines = []
        try:
            tabs = page.find_tables()
            if tabs and tabs.tables:
                for tab in tabs.tables:
                    extracted = tab.extract()
                    for r in extracted:
                        cleaned_r = [" ".join(str(c).split()) if c is not None else "" for c in r]
                        if any(cleaned_r):
                            table_lines.append("| " + " | ".join(cleaned_r) + " |")
        except Exception:
            pass

        # 2. Trích xuất văn bản thông thường
        text = page.get_text('text')
        if text:
            full_text.append(text.strip())
        if table_lines:
            full_text.append(f"\n[BẢNG BIỂU TRANG {page_idx}]:\n" + "\n".join(table_lines))
            
    doc.close()
    return '\n\n'.join(full_text)

def extract_file_content(filename: str, file_bytes: bytes) -> str:
    lower = filename.lower()
    if lower.endswith('.docx') or lower.endswith('.doc'):
        return extract_text_from_docx(file_bytes)
    elif lower.endswith('.pdf'):
        return extract_text_from_pdf(file_bytes)
    elif lower.endswith('.txt'):
        return file_bytes.decode('utf-8', errors='replace')
    else:
        raise ValueError('Định dạng tệp không hợp lệ. Vui lòng tải lên file docx hoặc pdf')

